"""Wave 2 W2-C: application/packages 用例测试(包导出/导入/Excel)。

覆盖要求: 用例可独立执行并提交事务(新会话可见, 无需调用方 commit)。
（旧 services.package 已删除，等价对照测试已随之删除。）

环境: SQLite :memory:(StaticPool) + 内存队列 + 临时对象存储目录。
"""

from __future__ import annotations

import io
import json
import os
import zipfile
from collections.abc import Iterator
from pathlib import Path

# 单文件运行时的安全网: 固定 SQLite + 内存队列, 避免误连部署 Postgres/Redis
os.environ.setdefault("IESPLAN_DB_URL", "sqlite+pysqlite://")
os.environ.setdefault("IESPLAN_QUEUE", "memory")

import pytest  # noqa: E402
from auth_helpers import login_headers, make_user  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from sqlalchemy import create_engine  # noqa: E402
from sqlalchemy.engine import Engine  # noqa: E402
from sqlalchemy.orm import Session, sessionmaker  # noqa: E402
from sqlalchemy.pool import StaticPool  # noqa: E402

from iesplan import identity as identity_domain  # noqa: E402
from iesplan import package as package_domain  # noqa: E402
from iesplan import project as project_domain  # noqa: E402
from iesplan.api import projects as projects_api  # noqa: E402
from iesplan.application import packages as packages_uc  # noqa: E402
from iesplan.application.packages import transfers as transfers_uc  # noqa: E402
from iesplan.config import settings  # noqa: E402
from iesplan.db import Base, get_db  # noqa: E402
from iesplan.main import create_app  # noqa: E402
from iesplan.tasks.persistence import CalcSnapshot, Task  # noqa: E402
from iesplan.results.persistence import EvidencePackage, ResultAssessment, ResultIndex  # noqa: E402
from iesplan.storage import get_object, put_object  # noqa: E402

# ---------------------------------------------------------------------------
# 测试环境
# ---------------------------------------------------------------------------


@pytest.fixture(scope="session")
def engine() -> Iterator[Engine]:
    eng = create_engine(
        "sqlite+pysqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(eng)
    yield eng
    eng.dispose()


@pytest.fixture(autouse=True)
def _clean_state(engine: Engine, db: Session) -> Iterator[None]:
    # 内存队列由 IESPLAN_QUEUE=memory 固定(模块顶部)，包流程不触队列，无需重置。
    yield
    with engine.begin() as conn:
        for table in reversed(Base.metadata.sorted_tables):
            conn.execute(table.delete())


@pytest.fixture()
def db(engine: Engine) -> Iterator[Session]:
    factory = sessionmaker(bind=engine, autoflush=False, expire_on_commit=False)
    with factory() as session:
        yield session


@pytest.fixture()
def client(engine: Engine, db: Session, tmp_path: Path) -> Iterator[TestClient]:
    settings.data_dir = tmp_path
    app = create_app()
    app.include_router(projects_api.router)

    def _override_get_db() -> Iterator[Session]:
        yield db

    app.dependency_overrides[get_db] = _override_get_db
    with TestClient(app, raise_server_exceptions=False) as test_client:
        yield test_client


def _h(client: TestClient, user) -> dict[str, str]:
    return login_headers(client, user)


def _user(db: Session, name: str):
    ns = make_user(db, name)
    db.commit()
    rec = identity_domain.get_user(db, ns.id)
    assert rec is not None
    return rec


def _project(client: TestClient, user, name: str) -> int:
    resp = client.post(
        "/api/projects",
        json={
            "name": name,
            "baseline_resolution": "1h",
            "baseline_leap_year": False,
            "baseline_scenario_mode": "single",
        },
        headers=_h(client, user),
    )
    assert resp.status_code == 201, resp.text
    return resp.json()["project"]["id"]


def _version(client: TestClient, user, project_id: int) -> int:
    resp = client.post(
        f"/api/projects/{project_id}/versions",
        json={"name": "基准版本 v1", "reason": "manual_save"},
        headers=_h(client, user),
    )
    assert resp.status_code == 201, resp.text
    return resp.json()["version"]["id"]


def _fresh_db(engine: Engine) -> Session:
    return sessionmaker(bind=engine, autoflush=False, expire_on_commit=False)()


def _seed_evidence(
    db: Session, project_id: int, version_id: int, owner_id: int
) -> tuple[int, int]:
    """直接建任务/快照/证据包/评估/结果索引, 返回 (evidence_package_id, assessment_id)。"""
    content = json.dumps(
        {
            "kpis": [{"name": "年购电量", "value": 123456.7, "unit": "kWh"}],
            "financial": {"总投资": 1000, "IRR": 0.12},
            "environmental": {"年碳排放": 123.4},
            "engineering": {"装机容量": 500},
            "applicability": {"适用范围": "华东", "限制": "无"},
        },
        ensure_ascii=False,
    ).encode("utf-8")
    obj = put_object(db, content, "application/json", source_category="evidence")
    snapshot = CalcSnapshot(
        project_version_id=version_id, dataset_version_ids=[],
        calc_config_snapshot={"params": {"horizon_years": 1}, "algorithm": "milp"},
        program_version="0.1.0", extension_versions={}, random_seed=42,
        tolerances={}, created_by=owner_id,
    )
    db.add(snapshot)
    db.flush()
    task = Task(project_id=project_id, type="calc", status="completed",
                business_outcome="normal_completion", calc_snapshot_id=snapshot.id,
                requested_by=owner_id)
    db.add(task)
    db.flush()
    package = EvidencePackage(
        task_id=task.id, calc_snapshot_id=snapshot.id, object_id=obj.id,
        status="complete", created_by=owner_id,
    )
    db.add(package)
    db.flush()
    assessment = ResultAssessment(
        evidence_package_id=package.id, assessor="system", assessed_by=owner_id,
        dimension_physical="pass", dimension_optimality="pass",
        dimension_financial="pass", dimension_reliability="pass",
        overall_score=88.5, comment="ok", detail={},
    )
    db.add(assessment)
    db.flush()
    db.add(ResultIndex(
        project_id=project_id, project_version_id=version_id,
        evidence_package_id=package.id, assessment_id=assessment.id, is_latest=True,
    ))
    db.commit()
    return package.id, assessment.id


# ---------------------------------------------------------------------------
# 包导出: 独立提交
# ---------------------------------------------------------------------------


def test_export_package_usecase_commits(
    client: TestClient, db: Session, engine: Engine
) -> None:
    owner = _user(db, "w2c_pkg_owner")
    pid = _project(client, owner, "w2c-导出项目")
    _version(client, owner, pid)
    user = identity_domain.get_user(db, owner.id)
    assert user is not None

    result = packages_uc.export_package(db, user, pid)
    assert result.manifest["format_version"] == "1.0"
    assert result.token

    # 事务提交证明: 新会话可读出包对象
    with _fresh_db(engine) as fresh:
        raw = get_object(fresh, result.object_id)
        assert raw[:2] == b"PK"
        manifest = json.loads(zipfile.ZipFile(io.BytesIO(raw)).read("manifest.json"))
        assert manifest["package_type"] == "project"

    assert result.file_name.endswith(".zip")


# ---------------------------------------------------------------------------
# 包导入: 提案幂等 + 确认建新身份
# ---------------------------------------------------------------------------


def test_import_roundtrip_via_usecases(client: TestClient, db: Session, engine: Engine) -> None:
    owner = _user(db, "w2c_pkg_owner2")
    pid = _project(client, owner, "w2c-导出项目2")
    _version(client, owner, pid)
    user = identity_domain.get_user(db, owner.id)
    assert user is not None
    exported = packages_uc.export_package(db, user, pid)
    zip_bytes = get_object(db, exported.object_id)

    importer = _user(db, "w2c_importer")
    imp = identity_domain.get_user(db, importer.id)
    assert imp is not None

    proposal = packages_uc.propose_import(db, imp, zip_bytes, idempotency_key="w2c-imp-1")
    assert proposal.status == "proposed"
    # 幂等: 同键返回同一提案
    again = packages_uc.propose_import(db, imp, zip_bytes, idempotency_key="w2c-imp-1")
    assert again.id == proposal.id

    new_project = packages_uc.confirm_import(db, imp, proposal.id)
    assert new_project.owner_id == imp.id
    assert new_project.id != pid
    # 事务提交证明: 新会话可见新项目
    with _fresh_db(engine) as fresh:
        got = project_domain.get_project(fresh, new_project.id)
        assert got is not None and got.owner_id == imp.id
    # 确认幂等重放: 已 applied 返回同一项目
    same = packages_uc.confirm_import(db, imp, proposal.id)
    assert same.id == new_project.id


def test_import_rejects_bad_package(client: TestClient, db: Session) -> None:
    owner = _user(db, "w2c_pkg_owner3")
    user = identity_domain.get_user(db, owner.id)
    assert user is not None

    with pytest.raises(package_domain.ImportValidationError) as e1:
        packages_uc.propose_import(db, user, b"not a zip at all")
    assert e1.value.code == "PKG-IMP-001"


# ---------------------------------------------------------------------------
# Excel 导出用例
# ---------------------------------------------------------------------------


def test_export_excel_usecase(client: TestClient, db: Session) -> None:
    owner = _user(db, "w2c_xls_owner")
    pid = _project(client, owner, "w2c-Excel项目")
    vid = _version(client, owner, pid)
    ep_id, a_id = _seed_evidence(db, pid, vid, owner.id)
    user = identity_domain.get_user(db, owner.id)
    assert user is not None

    new_bytes = transfers_uc.export_excel(db, user, pid, ep_id, a_id, lang="zh")
    assert new_bytes[:2] == b"PK"

    new_title = load_workbook(io.BytesIO(new_bytes))["报告总览"].cell(row=1, column=1).value
    assert "pIES 项目结果报告" in new_title


def test_download_token_helpers_match_old() -> None:
    token = packages_uc.create_download_token(7, "package", project_id=3, user_id=5)
    parsed = packages_uc.verify_download_token(token, expected_kind="package")
    assert parsed == {"object_id": 7, "kind": "package", "project_id": 3, "user_id": 5}
    old_parsed = package_domain.verify_download_token(token, expected_kind="package")
    assert old_parsed == parsed
